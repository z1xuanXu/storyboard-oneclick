import argparse
import json
from copy import copy
from pathlib import Path

from openpyxl import Workbook


def yunwu_node(node_id, order, x, y, title, prompt, cfg):
    return {
        "id": node_id,
        "type": "YunwuGeminiImage",
        "pos": [x, y],
        "size": [520, 1029.3125],
        "flags": {},
        "order": order,
        "mode": 0,
        "inputs": [{"name": f"image_{i}", "type": "IMAGE", "link": None} for i in range(1, 15)],
        "outputs": [
            {"name": "images", "type": "IMAGE", "slot_index": 0, "links": []},
            {"name": "text", "type": "STRING", "slot_index": 1, "links": []},
            {"name": "raw_json", "type": "STRING", "slot_index": 2, "links": []},
        ],
        "title": title,
        "properties": {"Node name for S&R": "YunwuGeminiImage"},
        "widgets_values": [
            prompt,
            "",
            cfg["base_url"],
            cfg["model"],
            "TEXT,IMAGE",
            cfg["aspect_ratio"],
            cfg["image_size"],
            1,
            1,
            0,
            1,
            0,
            -1,
            "randomize",
            "image/png",
            "{}",
            "{}",
        ],
    }


def save_node(node_id, order, x, y, title, prefix):
    return {
        "id": node_id,
        "type": "SaveImage",
        "pos": [x, y],
        "size": [320, 84],
        "flags": {},
        "order": order,
        "mode": 0,
        "inputs": [{"name": "images", "type": "IMAGE", "link": None}],
        "outputs": [],
        "title": title,
        "properties": {"Node name for S&R": "SaveImage"},
        "widgets_values": [prefix],
    }


def build_workbook(data, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "项目说明"
    ws.append(["项目", data["title"]])
    ws.append(["比例", data["aspect_ratio"]])
    ws.append(["尺寸", data["image_size"]])
    ws.append(["模型", data["model"]])

    ref_ws = wb.create_sheet("参考图清单")
    ref_ws.append(["编号", "名称", "用途", "提示词", "用于镜头"])
    for ref in data["refs"]:
        ref_ws.append([ref["id"], ref["name"], ref["purpose"], ref["prompt"], ref["use_in_shots"]])

    shot_ws = wb.create_sheet("分镜脚本")
    shot_ws.append(["镜号", "场次", "建议时长", "剪辑/景别", "画面详细描述", "单帧图提示词", "参考图输入", "图生视频提示词", "对白/音效", "生成状态", "单帧图粘贴区"])
    for shot in data["shots"]:
        shot_ws.append([
            shot["id"],
            shot["scene"],
            shot["duration"],
            shot["cut"],
            shot["visual"],
            shot["frame_prompt"],
            ", ".join(shot["references"]),
            shot["video_prompt"],
            shot["dialogue_sound"],
            "待生成",
            "",
        ])

    run_ws = wb.create_sheet("ComfyUI执行清单")
    run_ws.append(["编号", "类型", "说明"])
    for ref in data["refs"]:
        run_ws.append([ref["id"], "参考图", ref["purpose"]])
    for shot in data["shots"]:
        run_ws.append([shot["id"], "镜头图", "使用参考图：" + ", ".join(shot["references"])])

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "top"
                cell.alignment = alignment
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[chr(64 + column)].width = 18 if column < 5 else 36
    shot_ws.column_dimensions["K"].width = 38
    wb.save(out_path)


def build_workflow(data, out_path):
    prompts = []
    for ref in data["refs"]:
        prompts.append({"id": ref["id"], "title": ref["name"], "prompt": ref["prompt"], "refs": []})
    for shot in data["shots"]:
        prompts.append({"id": shot["id"], "title": shot["scene"], "prompt": shot["frame_prompt"], "refs": shot["references"]})

    nodes = []
    node_by_prompt = {}
    next_id = 1
    for index, item in enumerate(prompts):
        col = index % 4
        row = index // 4
        node = yunwu_node(next_id, index, col * 620, row * 860, f"{item['id']} {item['title']}", item["prompt"], data)
        nodes.append(node)
        node_by_prompt[item["id"]] = node
        next_id += 1

    links = []
    next_link = 1
    for item in prompts:
        target = node_by_prompt[item["id"]]
        for input_index, ref_id in enumerate(item["refs"]):
            source = node_by_prompt.get(ref_id)
            if not source:
                continue
            links.append([next_link, source["id"], 0, target["id"], input_index, "IMAGE"])
            source["outputs"][0]["links"].append(next_link)
            target["inputs"][input_index]["link"] = next_link
            next_link += 1

    for item in prompts:
        source = node_by_prompt[item["id"]]
        save = save_node(next_id, len(nodes), source["pos"][0], source["pos"][1] + 780, f"Save {item['id']}", f"phone_key_story/{item['id']}")
        nodes.append(save)
        links.append([next_link, source["id"], 0, save["id"], 0, "IMAGE"])
        source["outputs"][0]["links"].append(next_link)
        save["inputs"][0]["link"] = next_link
        next_id += 1
        next_link += 1

    workflow = {
        "id": "phone-key-story-yunwu-template",
        "revision": 0,
        "last_node_id": len(nodes),
        "last_link_id": len(links),
        "nodes": nodes,
        "links": links,
        "groups": [
            {"id": 1, "title": "Generate references first", "bounding": [-40, -70, 2500, 900], "color": "#0F766E", "font_size": 24, "flags": {}},
            {"id": 2, "title": "Shot still generation", "bounding": [-40, 790, 2500, 5400], "color": "#7C2D12", "font_size": 24, "flags": {}},
        ],
        "config": {},
        "extra": {"ds": {"scale": 0.55, "offset": [80, 80]}},
        "version": 0.4,
    }
    out_path.write_text(json.dumps(workflow, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--story", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    data = json.loads(args.story.read_text(encoding="utf-8"))
    args.out.mkdir(parents=True, exist_ok=True)
    build_workbook(data, args.out / "phone_key_storyboard.xlsx")
    build_workflow(data, args.out / "phone_key_story_yunwu_template.json")
    print({"workbook": str(args.out / "phone_key_storyboard.xlsx"), "workflow": str(args.out / "phone_key_story_yunwu_template.json")})


if __name__ == "__main__":
    main()
