import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XlsxImage


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--images", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    wb = load_workbook(args.workbook)
    ws = wb["分镜脚本"]
    ws.column_dimensions["K"].width = 38
    ws.column_dimensions["J"].width = 18

    missing = []
    inserted = 0
    for row in range(2, ws.max_row + 1):
        shot_id = ws[f"A{row}"].value
        if not isinstance(shot_id, str) or not shot_id.startswith("S"):
            continue
        candidates = sorted(args.images.glob(f"{shot_id}_*.png"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not candidates:
            missing.append(shot_id)
            ws[f"J{row}"] = "未找到输出图"
            continue
        image = XlsxImage(str(candidates[0]))
        image.width = 250
        image.height = 141
        image.anchor = f"K{row}"
        ws.add_image(image)
        ws.row_dimensions[row].height = 118
        ws[f"J{row}"] = "已生成并贴入"
        ws[f"K{row}"] = ""
        inserted += 1

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print({"output": str(args.out), "inserted": inserted, "missing": missing})


if __name__ == "__main__":
    main()

